import csv
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

from .extensions import db
from .models import Record, now_ist


class ParseError(Exception):
    """Raised for unreadable files or missing required columns."""


# field -> accepted normalised header names (lowercase, letters/digits only)
ALIASES = {
    "branch_name": ["branchname", "branch"],
    "member_code": ["membercode", "membercod", "membercd", "memberid"],
    "member_name": ["membername", "name"],
    "member_mobile": ["membermobilenumber", "membermobileno", "membermobile", "mobilenumber", "mobileno"],
    "account_number": ["accountnumber", "accountno", "accno", "account"],
    "open_date": ["opendate", "opendt"],
    "close_date": ["closedate", "closedt"],
    "due": ["due", "dues"],
    "policy_amt": ["policyamt", "policyamount"],
    "amount_due": ["amountdue"],
    "dist_status": ["diststatus", "distributionstatus"],
    "delivered_to": ["deliveredto", "diliveredto"],
    "mob": ["mob"],
    "kyc_status": ["kycstatus", "kyc"],
    "remark": ["remark", "remarks"],
}
REQUIRED = ["member_code", "member_name"]
TEXT_FIELDS = [
    "branch_name", "member_code", "member_name", "member_mobile", "account_number",
    "open_date", "close_date", "dist_status", "delivered_to", "mob", "kyc_status", "remark",
]


def _norm(header):
    return re.sub(r"[^a-z0-9]", "", str(header).lower())


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        if value.is_integer():
            return str(int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}( 00:00:00)?", text):  # ISO date string
        y, m, d = text[:10].split("-")
        return f"{d}/{m}/{y}"
    return text


def _number(value):
    """Return float, or None when the cell is blank/unreadable."""
    text = _text(value).replace(",", "").replace("₹", "").replace("Rs.", "").strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _dist_status(value):
    text = value.lower()
    if not text:
        return ""
    if "not" in text or "pend" in text:
        return "Pending"
    if any(word in text for word in ("deliver", "delivr", "dilivered", "done", "given", "taken", "yes")):
        return "Delivered"
    return "Pending"


def _read_table(file_storage):
    """Return (header_row, data_rows) from an .xlsx/.xlsm/.csv upload (no pandas/numpy needed)."""
    filename = (file_storage.filename or "").lower()
    raw = file_storage.read()
    try:
        if filename.endswith(".csv"):
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("cp1252")
            rows = list(csv.reader(io.StringIO(text)))
        elif filename.endswith((".xlsx", ".xlsm")):
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                sheet.reset_dimensions()  # some exports store a wrong sheet size
                rows = [list(r) for r in sheet.iter_rows(values_only=True)]
            finally:
                workbook.close()
        else:
            raise ParseError("Please upload an .xlsx or .csv file.")
    except ParseError:
        raise
    except Exception as exc:  # corrupt file, wrong format, etc.
        raise ParseError(f"Could not read the file: {exc}")

    for index, row in enumerate(rows):
        if any(_text(cell) for cell in row):
            return row, rows[index + 1:]
    raise ParseError("The file is empty.")


def parse_file(file_storage):
    """Read an uploaded .xlsx/.csv. Returns (rows, skipped_count)."""
    header, data_rows = _read_table(file_storage)

    found = {}  # field -> column index
    for index, cell in enumerate(header):
        key = _norm(cell) if cell is not None else ""
        for field, aliases in ALIASES.items():
            if key in aliases and field not in found:
                found[field] = index
    missing = [f for f in REQUIRED if f not in found]
    if missing:
        raise ParseError(
            "Missing required column(s): "
            + ", ".join(m.replace("_", " ").title() for m in missing)
            + ". Check the header row (first row) of your sheet."
        )

    def cell(source, field):
        index = found.get(field)
        return source[index] if index is not None and index < len(source) else None

    rows, skipped = [], 0
    for source in data_rows:
        if not any(_text(c) for c in source):
            continue  # fully blank line
        row = {f: _text(cell(source, f)) for f in TEXT_FIELDS}
        if not row["member_code"] or not row["member_name"]:
            skipped += 1
            continue
        due = _number(cell(source, "due"))
        row["due"] = max(int(round(due)), 0) if due is not None else 0
        policy = _number(cell(source, "policy_amt"))
        row["policy_amt"] = max(policy, 0) if policy is not None else None
        row["dist_status"] = _dist_status(row["dist_status"])
        rows.append(row)
    return rows, skipped


def import_rows(rows, can_touch=None):
    """Insert new records / update existing ones (key: member code + account no).

    Master fields always come from the file; tracking fields (distribution status,
    delivered to, mob, KYC, remark, policy amount) only overwrite when the file
    has a value, so work done in the app is not lost on re-upload.
    """
    existing = {(r.member_code, r.account_number): r for r in Record.query.all()}
    added = updated = blocked = 0
    for data in rows:
        key = (data["member_code"], data["account_number"])
        record = existing.get(key)
        if record is not None and can_touch and not can_touch(record.branch_name):
            blocked += 1  # row belongs to another branch
            continue
        is_new = record is None
        if is_new:
            record = Record(
                member_code=data["member_code"],
                account_number=data["account_number"],
                dist_status="Pending",
                policy_amt=0,
            )
            db.session.add(record)
            existing[key] = record
            added += 1
        else:
            updated += 1

        record.branch_name = data["branch_name"]
        record.member_name = data["member_name"]
        record.member_mobile = data["member_mobile"]
        record.open_date = data["open_date"]
        record.close_date = data["close_date"]
        record.due = data["due"]
        if data["policy_amt"] is not None:
            record.policy_amt = data["policy_amt"]
        for field in ("delivered_to", "mob", "kyc_status"):
            if data[field]:
                setattr(record, field, data[field])
        if data["dist_status"]:
            if data["dist_status"] == "Delivered" and record.dist_status != "Delivered":
                record.delivered_at = now_ist()
            record.dist_status = data["dist_status"]
        record.recalc()
        # keep a remark from the file only if it differs from the auto-generated one
        if data["remark"] and data["remark"] != record.auto_remark():
            record.remark = data["remark"]
            record.remark_manual = True
    db.session.commit()
    return added, updated, blocked


def preview_rows(rows, can_touch=None):
    """Count what import_rows would do, without changing anything."""
    known = {
        (code, account): branch
        for code, account, branch in Record.query.with_entities(
            Record.member_code, Record.account_number, Record.branch_name
        )
    }
    new = update = blocked = 0
    for data in rows:
        key = (data["member_code"], data["account_number"])
        if key in known:
            if can_touch and not can_touch(known[key]):
                blocked += 1
            else:
                update += 1
        else:
            new += 1
            known[key] = data["branch_name"]
    return {"new": new, "update": update, "blocked": blocked}